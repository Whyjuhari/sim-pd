from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCUMENTS = ROOT / "resources" / "documents"
TEMPLATES = {
    "Template_SPT_Dasar_Regulasi.docx":
        "Template_SPT_Dasar_Regulasi_Lampiran.docx",
    "Template_SPT_Dasar_Regulasi_Memo.docx":
        "Template_SPT_Dasar_Regulasi_Memo_Lampiran.docx",
    "Template_SPT_Dasar_Regulasi_DIPA.docx":
        "Template_SPT_Dasar_Regulasi_DIPA_Lampiran.docx",
    "Template_SPT_Dasar_Regulasi_Memo_DIPA.docx":
        "Template_SPT_Dasar_Regulasi_Memo_DIPA_Lampiran.docx",
}

# The travel table must fit inside the printable height of the existing A4
# section. These values are minimum heights (twips), not fixed heights, so Word
# can still grow a row if its real content needs more room.
COMPACT_TRAVEL_ROW_HEIGHTS = (2750, 2400, 2550, 2450, 3070, 234, 650, 850)


def set_font(run, size: float = 11, bold: bool = False) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:eastAsia"), "Arial")
    fonts.set(qn("w:cs"), "Arial")


def format_paragraph(
    paragraph,
    *,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    keep_next: bool = False,
) -> None:
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1

    paragraph_properties = paragraph._p.get_or_add_pPr()
    keep_lines = paragraph_properties.find(qn("w:keepLines"))
    if keep_lines is None:
        paragraph_properties.append(OxmlElement("w:keepLines"))

    if keep_next and paragraph_properties.find(qn("w:keepNext")) is None:
        paragraph_properties.append(OxmlElement("w:keepNext"))


def clear_cell(cell) -> None:
    cell_properties = cell._tc.tcPr
    for child in list(cell._tc):
        if child is not cell_properties:
            cell._tc.remove(child)


def set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    size: float = 11,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    clear_cell(cell)
    paragraph = cell.add_paragraph()
    format_paragraph(paragraph, alignment=alignment)
    set_font(paragraph.add_run(text), size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_width(cell, width_cm: float) -> None:
    width = Cm(width_cm)
    cell.width = width
    cell_properties = cell._tc.get_or_add_tcPr()
    cell_width = cell_properties.find(qn("w:tcW"))
    if cell_width is None:
        cell_width = OxmlElement("w:tcW")
        cell_properties.append(cell_width)
    cell_width.set(qn("w:w"), str(int(width.twips)))
    cell_width.set(qn("w:type"), "dxa")


def set_cell_margins(cell, value: int = 80) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)

    for edge in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders_none(cell) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    borders = cell_properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        cell_properties.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_cell_borders(
    cell,
    *,
    value: str = "single",
    size: int = 4,
    color: str = "000000",
) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    borders = cell_properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        cell_properties.append(borders)

    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), value)
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_borders(table, value: str) -> None:
    table_properties = table._tbl.tblPr
    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), value)
        if value != "nil":
            node.set(qn("w:sz"), "4")
            node.set(qn("w:space"), "0")
            node.set(qn("w:color"), "000000")


def set_table_layout_fixed(table) -> None:
    table_properties = table._tbl.tblPr
    layout = table_properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")


def keep_row_together(row, repeat_header: bool = False) -> None:
    # PHPWord prioritizes row start tags that contain attributes when cloning.
    # Give generated rows an rsid so cloneRowAndSetValues selects this row,
    # not an earlier row from another table in the document.
    row._tr.set(qn("w:rsidR"), "00000000")
    row_properties = row._tr.get_or_add_trPr()
    if row_properties.find(qn("w:cantSplit")) is None:
        row_properties.append(OxmlElement("w:cantSplit"))
    if repeat_header and row_properties.find(qn("w:tblHeader")) is None:
        row_properties.append(OxmlElement("w:tblHeader"))


def set_row_minimum_height(row, height_twips: int) -> None:
    row_properties = row._tr.get_or_add_trPr()
    heights = row_properties.findall(qn("w:trHeight"))
    height = heights[-1] if heights else OxmlElement("w:trHeight")

    for duplicate in heights[:-1]:
        row_properties.remove(duplicate)
    if not heights:
        row_properties.append(height)

    height.set(qn("w:val"), str(height_twips))
    height.set(qn("w:hRule"), "atLeast")


def remove_paragraph(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def trim_sender_signature_spacing_in_row(first_row) -> None:
    """Remove only redundant blank lines around the row-I signature token."""
    signature_cell = next(
        (cell for cell in first_row.cells if "${ttd_pengirim}" in cell.text),
        None,
    )
    if signature_cell is None:
        return

    paragraphs = list(signature_cell.paragraphs)
    signature_index = next(
        index
        for index, paragraph in enumerate(paragraphs)
        if "${ttd_pengirim}" in paragraph.text
    )

    blanks_before = []
    cursor = signature_index - 1
    while cursor >= 0 and not paragraphs[cursor].text.strip():
        blanks_before.append(paragraphs[cursor])
        cursor -= 1

    blanks_after = []
    cursor = signature_index + 1
    while cursor < len(paragraphs) and not paragraphs[cursor].text.strip():
        blanks_after.append(paragraphs[cursor])
        cursor += 1

    # Keep one line above and two below the signature token. This retains the
    # intended signature area while recovering enough height for rows VI-VII.
    for paragraph in blanks_before[1:]:
        remove_paragraph(paragraph)
    for paragraph in blanks_after[2:]:
        remove_paragraph(paragraph)


def trim_sender_signature_spacing(table) -> None:
    trim_sender_signature_spacing_in_row(table.rows[0])


def find_travel_table(document):
    for table in reversed(document.tables):
        text = "\n".join(cell.text for row in table.rows for cell in row.cells)
        if "VI. Catatan lain-lain" in text and "VII. PERHATIAN" in text:
            return table
    raise RuntimeError("Tabel perjalanan I-VII tidak ditemukan.")


def remove_empty_paragraphs_after(table) -> None:
    sibling = table._tbl.getnext()
    while sibling is not None and sibling.tag == qn("w:p"):
        following = sibling.getnext()
        has_content = any(
            sibling.find(f".//{qn(tag)}") is not None
            for tag in ("w:t", "w:br", "w:drawing", "w:pict", "w:sectPr")
        )
        if has_content:
            break
        sibling.getparent().remove(sibling)
        sibling = following


def keep_paragraph_chain(paragraphs) -> None:
    for index, paragraph in enumerate(paragraphs):
        properties = paragraph.get_or_add_pPr()
        if properties.find(qn("w:keepLines")) is None:
            properties.append(OxmlElement("w:keepLines"))

        keep_next = properties.find(qn("w:keepNext"))
        if index < len(paragraphs) - 1:
            if keep_next is None:
                properties.append(OxmlElement("w:keepNext"))
        elif keep_next is not None:
            properties.remove(keep_next)


def compact_travel_rows(table, rows) -> None:
    if len(rows) != len(COMPACT_TRAVEL_ROW_HEIGHTS):
        raise RuntimeError(
            f"Tabel perjalanan memiliki {len(rows)} baris; "
            f"seharusnya {len(COMPACT_TRAVEL_ROW_HEIGHTS)}."
        )

    trim_sender_signature_spacing_in_row(rows[0])
    remove_empty_paragraphs_after(table)

    for row, height_twips in zip(rows, COMPACT_TRAVEL_ROW_HEIGHTS):
        keep_row_together(row)
        set_row_minimum_height(row, height_twips)

    paragraphs = []
    for row in rows:
        paragraphs.extend(row._tr.iter(qn("w:p")))
    keep_paragraph_chain(paragraphs)


def compact_travel_table(document) -> None:
    table = find_travel_table(document)
    compact_travel_rows(table, list(table.rows))


def page_break_element():
    paragraph = OxmlElement("w:p")
    properties = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    properties.append(spacing)
    paragraph.append(properties)
    run = OxmlElement("w:r")
    page_break = OxmlElement("w:br")
    page_break.set(qn("w:type"), "page")
    run.append(page_break)
    paragraph.append(run)
    return paragraph


def split_inline_closing_and_travel(table) -> None:
    closing_table = deepcopy(table._tbl)
    travel_table = deepcopy(table._tbl)
    keep_table_rows(closing_table, keep_first=True)
    keep_table_rows(travel_table, keep_first=False)

    table._tbl.addprevious(closing_table)
    table._tbl.addprevious(page_break_element())
    table._tbl.addprevious(travel_table)
    table._tbl.getparent().remove(table._tbl)


def compact_inline_travel_table(document) -> None:
    table = find_travel_table(document)
    if len(table.rows) == len(COMPACT_TRAVEL_ROW_HEIGHTS):
        compact_travel_rows(table, list(table.rows))
        return
    if len(table.rows) != len(COMPACT_TRAVEL_ROW_HEIGHTS) + 1:
        raise RuntimeError(
            "Template inline harus memiliki satu baris penutup dan delapan "
            "baris perjalanan."
        )
    compact_travel_rows(table, list(table.rows)[1:])
    split_inline_closing_and_travel(table)


def add_metadata(container) -> None:
    table = container.add_table(rows=3, cols=4)
    table.autofit = False
    set_table_layout_fixed(table)
    set_table_borders(table, "nil")
    widths = (9.5, 1.7, 0.4, 7.1)

    for row in table.rows:
        for column, width in enumerate(widths):
            set_cell_width(row.cells[column], width)
            set_cell_margins(row.cells[column], 0)
            set_cell_borders_none(row.cells[column])

    title_cell = table.cell(0, 1).merge(table.cell(0, 3))
    set_cell_text(title_cell, "Lampiran Surat Tugas", size=11)

    for row_index, (label, value) in enumerate((
        ("Nomor", "${nomor_surat}"),
        ("Tanggal", "${tanggal_surat}"),
    ), start=1):
        set_cell_text(table.cell(row_index, 0), "", size=11)
        set_cell_text(table.cell(row_index, 1), label, size=11)
        set_cell_text(table.cell(row_index, 2), ":", size=11)
        set_cell_text(table.cell(row_index, 3), value, size=11)


def add_employee_table(container) -> None:
    title = container.add_paragraph()
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(8)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_properties = title._p.get_or_add_pPr()
    title_properties.append(OxmlElement("w:keepNext"))
    title_properties.append(OxmlElement("w:keepLines"))
    set_font(
        title.add_run("DAFTAR PEJABAT/PEGAWAI YANG DIBERI TUGAS"),
        size=12,
        bold=True,
    )

    table = container.add_table(rows=2, cols=5)
    table.autofit = False
    set_table_layout_fixed(table)
    set_table_borders(table, "single")
    widths = (1.0, 4.5, 4.2, 4.0, 5.0)
    headers = ("NO", "NAMA", "NIP/NIK", "PANGKAT/GOL.\nRUANG", "JABATAN")
    values = (
        "${nomor_pegawai}",
        "${nama_pegawai}",
        "${nip}",
        "${pangkat_golongan}",
        "${jabatan}",
    )

    for column, width in enumerate(widths):
        set_cell_width(table.cell(0, column), width)
        set_cell_width(table.cell(1, column), width)
        set_cell_margins(table.cell(0, column), 100)
        set_cell_margins(table.cell(1, column), 100)
        set_cell_borders(table.cell(0, column))
        set_cell_borders(table.cell(1, column))
        set_cell_text(
            table.cell(0, column),
            headers[column],
            bold=True,
            size=10.5,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_text(
            table.cell(1, column),
            values[column],
            size=10.5,
            alignment=(
                WD_ALIGN_PARAGRAPH.CENTER
                if column in (0, 2, 3)
                else WD_ALIGN_PARAGRAPH.LEFT
            ),
        )

    keep_row_together(table.rows[0], repeat_header=True)
    keep_row_together(table.rows[1])


def add_signature(container) -> None:
    table = container.add_table(rows=1, cols=2)
    table.autofit = False
    set_table_borders(table, "nil")
    set_cell_width(table.cell(0, 0), 10.5)
    set_cell_width(table.cell(0, 1), 8.0)
    set_cell_borders_none(table.cell(0, 0))
    set_cell_borders_none(table.cell(0, 1))
    clear_cell(table.cell(0, 0))
    table.cell(0, 0).add_paragraph()

    right = table.cell(0, 1)
    clear_cell(right)
    signature_lines = (
        ("Kepala Balai,", False, 20),
        ("${ttd_pengirim}", False, 42),
        ("Ashari Arifuddin, S.T., M.M", False, 20),
        ("NIP. 197606242011011001", False, 0),
    )
    for index, (text, bold, spacing_before) in enumerate(signature_lines):
        paragraph = right.add_paragraph()
        format_paragraph(
            paragraph,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            keep_next=index < len(signature_lines) - 1,
        )
        paragraph.paragraph_format.space_before = Pt(spacing_before)
        set_font(paragraph.add_run(text), size=11, bold=bold)


def keep_table_rows(table_element, *, keep_first: bool) -> None:
    rows = list(table_element.findall(qn("w:tr")))
    selected_rows = rows[:1] if keep_first else rows[1:]

    for row in rows:
        if row not in selected_rows:
            table_element.remove(row)


def make_attachment_section(document) -> None:
    original_table = document.tables[1]._tbl
    closing_table = deepcopy(original_table)
    travel_table = deepcopy(original_table)
    keep_table_rows(closing_table, keep_first=True)
    keep_table_rows(travel_table, keep_first=False)

    generated_elements = []

    start = document.add_paragraph()
    format_paragraph(start, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    start.add_run().add_break(WD_BREAK.PAGE)
    set_font(start.add_run("- 2 -"), size=11)
    generated_elements.append(start._p)

    before_metadata = len(document.tables)
    add_metadata(document)
    generated_elements.append(document.tables[before_metadata]._tbl)

    before_title = len(document.paragraphs)
    before_employee_table = len(document.tables)
    add_employee_table(document)
    generated_elements.append(document.paragraphs[before_title]._p)
    generated_elements.append(document.tables[before_employee_table]._tbl)

    before_signature = len(document.tables)
    add_signature(document)
    generated_elements.append(document.tables[before_signature]._tbl)

    end = document.add_paragraph()
    format_paragraph(end)
    end.add_run().add_break(WD_BREAK.PAGE)
    generated_elements.append(end._p)

    original_table.addprevious(closing_table)
    for element in generated_elements:
        original_table.addprevious(element)
    original_table.addprevious(travel_table)
    original_table.getparent().remove(original_table)


def replace_main_employee_row(document) -> None:
    row = document.tables[0].rows[3]
    set_cell_text(row.cells[0], "Kepada", size=12)
    set_cell_text(row.cells[1], ":", size=12)
    employee_cell = row.cells[2].merge(row.cells[4])
    set_cell_text(employee_cell, "Nama-nama Terlampir", size=12)


def build(source_name: str, output_name: str) -> None:
    source = DOCUMENTS / source_name
    output = DOCUMENTS / output_name
    document = Document(source)
    replace_main_employee_row(document)
    make_attachment_section(document)
    compact_travel_table(document)
    document.save(output)


def compact_existing(output_name: str) -> None:
    output = DOCUMENTS / output_name
    document = Document(output)
    compact_travel_table(document)
    document.save(output)


def compact_inline_existing(source_name: str) -> None:
    source = DOCUMENTS / source_name
    document = Document(source)
    compact_inline_travel_table(document)
    document.save(source)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--compact-existing",
        action="store_true",
        help="Padatkan hanya tabel perjalanan pada file lampiran yang ada.",
    )
    parser.add_argument(
        "--compact-inline",
        action="store_true",
        help="Padatkan tabel perjalanan pada template satu-dua pegawai.",
    )
    arguments = parser.parse_args()

    for source_name, output_name in TEMPLATES.items():
        if arguments.compact_inline:
            compact_inline_existing(source_name)
        elif arguments.compact_existing:
            compact_existing(output_name)
        else:
            build(source_name, output_name)
