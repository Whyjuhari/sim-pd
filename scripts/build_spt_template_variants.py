from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "resources" / "documents" / "Template_SPT_Perjadin_COLLECTIVE_V5.docx"
OUTPUTS = {
    "Template_SPT_Dasar_Regulasi.docx": (False, False),
    "Template_SPT_Dasar_Regulasi_Memo.docx": (True, False),
    "Template_SPT_Dasar_Regulasi_DIPA.docx": (False, True),
    "Template_SPT_Dasar_Regulasi_Memo_DIPA.docx": (True, True),
}


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def replace_paragraph_text(paragraph, text: str) -> None:
    first_run = paragraph.runs[0]
    first_run.text = text
    for run in paragraph.runs[1:]:
        run._element.getparent().remove(run._element)


def replace_runs(paragraph, replacements: dict[str, str]) -> None:
    for run in paragraph.runs:
        for old, new in replacements.items():
            if old in run.text:
                run.text = run.text.replace(old, new)


def build(output_name: str, uses_memo: bool, uses_dipa: bool) -> None:
    document = Document(SOURCE)
    basis_cell = document.tables[0].cell(1, 2)
    memo = basis_cell.paragraphs[3]
    dipa = basis_cell.paragraphs[4]

    if uses_dipa:
        replace_paragraph_text(
            dipa,
            "Biaya akibat surat tugas ini dibebankan pada Anggaran yang tertuang "
            "dalam DIPA BPVP Pangkajene dan Kepulauan TA. ${tahun_anggaran} "
            "Nomor : ${nomor_dipa}, tanggal ${tanggal_dipa}.",
        )
    else:
        remove_paragraph(dipa)

    if not uses_memo:
        remove_paragraph(memo)

    untuk_cell = document.tables[0].cell(5, 2)
    replace_runs(untuk_cell.paragraphs[2], {"TA. 2025": "TA. ${tahun_anggaran}"})

    output = SOURCE.parent / output_name
    document.save(output)


if __name__ == "__main__":
    for name, flags in OUTPUTS.items():
        build(name, *flags)
